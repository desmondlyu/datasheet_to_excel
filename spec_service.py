"""Normalize PDF geometry without filling empty limits from adjacent rows."""
from collections import Counter, defaultdict
from io import BytesIO
from pathlib import Path
import re
import pdfplumber
import fitz
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from extraction_core import extract_tables, ELECTRICAL_SECTION_LABELS, PDF_TABLE_SETTINGS

COLUMNS = ['產品名稱','max','typ','min','symbol','parameter','condition','description','unit','spec type','page']
PRODUCT_RE = re.compile(r'W25Q(?:32|64|12|25|51|01|02)RW(?:-DTR)?', re.I)
PRODUCT_ORDER = ['W25Q'+x+'RW' for x in ['32','64','12','25','51','01','02']]

def clean(v):
    return re.sub(r'\s+', ' ', str('' if v is None else v)).strip()

def products(text, defaults):
    # Abbreviated ranges in this family (e.g. W25Q32RW-01RW).
    m=re.search(r'W25Q(32|64|12|25|51|01|02)RW\s*[-–]\s*(?:W25Q)?(32|64|12|25|51|01|02)RW',text,re.I)
    if m:
        a,b=('W25Q'+x.upper()+'RW' for x in m.groups())
        return PRODUCT_ORDER[PRODUCT_ORDER.index(a):PRODUCT_ORDER.index(b)+1]
    found=list(dict.fromkeys(x.upper().removesuffix('-DTR') for x in PRODUCT_RE.findall(text)))
    return found or defaults

def value(v):
    v=clean(v)
    if re.fullmatch(r'[-−–]?\d[\d,]*(?:\.\d+)?',v):
        n=float(v.replace(',','').replace('−','-').replace('–','-'))
        return int(n) if n.is_integer() else n
    return v

def strip_product(text):
    return clean(re.sub(r'W25Q\w+(?:\s*[-–]\s*(?:W25Q)?\w+RW)?','',text)).strip(' ,;')

def row_record(product, vals, page, kind, section, notes):
    symbol=clean(vals.get('symbol',''))
    refs=re.findall(r'\((\d+)\)',symbol)
    symbol=clean(re.sub(r'\(\d+\)','',symbol))
    symbol=re.sub(r'\b([tVICF])\s+(?=[A-Z])',r'\1',symbol)
    r=dict.fromkeys(COLUMNS,'')
    r.update({'產品名稱':product,'symbol':symbol,'parameter':clean(vals.get('parameter')),
        'condition':strip_product(vals.get('condition','')),'description':'; '.join('Note '+n for n in refs),
        'unit':clean(vals.get('unit')),'spec type':kind,'page':page})
    for k in ['min','typ','max']:
        raw=clean(vals.get(k,''))
        limit_refs=re.findall(r'\((\d+)\)',raw)
        if limit_refs:
            r['description']+='; '+k+': '+', '.join('Note '+n for n in limit_refs)
            raw=re.sub(r'\(\d+\)','',raw)
        r[k]=value(raw)
    if 'Factory Mode' in section:
        r['condition']=('; '.join(filter(None,[r['condition'],'Factory Mode'])))
    r['_section']=section
    r['_notes']=notes
    return r

def cell_text(page, cell):
    return clean(page.crop(cell).extract_text(x_tolerance=2,y_tolerance=3))

def standard_table(page, table, defaults, page_no, kind, section, notes):
    raw=table.extract(); headers={}
    for row,geom in zip(raw[:3],table.rows[:3]):
        for txt,cell in zip(row,geom.cells):
            token=clean(txt).lower()
            key={'description':'parameter','parameter':'parameter','symbol':'symbol','conditions':'condition','condition':'condition','min':'min','typ':'typ','max':'max','unit':'unit'}.get(token)
            if key and cell:headers.setdefault(key,cell)
    if not {'parameter','symbol'}.issubset(headers) or not any(k in headers for k in ['min','typ','max']):return []
    ordered=sorted(headers,key=lambda k:headers[k][0])
    bounds={k:(headers[k][0],headers[ordered[i+1]][0] if i+1<len(ordered) else table.bbox[2]) for i,k in enumerate(ordered)}
    body_top=max(c[3] for c in headers.values())
    # Symbol cells define a whole parameter even when text has wrapped into artificial subrows.
    sx0,sx1=bounds['symbol']; anchors=[]
    for c in table.cells:
        if c[1]<body_top-1 or c[0]<sx0-1 or c[2]>sx1+1:continue
        txt=cell_text(page,c)
        if txt and re.search('[A-Za-z]',txt):anchors.append((c,txt))
    records=[]
    for anchor,sym in sorted(anchors,key=lambda a:(a[0][1],a[0][0])):
        y0,y1=anchor[1],anchor[3]
        px0,px1=bounds['parameter']
        param=cell_text(page,(px0,y0,px1,y1))
        # Genuine horizontal splits inside a parameter represent separate product/condition limits.
        ys={y0,y1}
        for c in table.cells:
            if c[1]>=y0-1 and c[3]<=y1+1 and c[0]>=bounds.get('condition',bounds.get('min',bounds['symbol']))[0]-1:
                ys.update([max(y0,c[1]),min(y1,c[3])])
        ys=sorted(ys); active_products=defaults
        for lo,hi in zip(ys,ys[1:]):
            if hi-lo<2:continue
            vals={'parameter':param,'symbol':sym}
            for key,(x0,x1) in bounds.items():
                if key in ['parameter','symbol']:continue
                matching=[c for c in table.cells if c[0]>=x0-1 and c[2]<=x1+1 and c[1] <= (lo+hi)/2 < c[3]]
                vals[key]=clean(' '.join(dict.fromkeys(cell_text(page,c) for c in sorted(matching))))
            if not any(vals.get(k) for k in ['min','typ','max']) and 'min' in bounds and 'max' in bounds:
                rx0=bounds['min'][0];rx1=bounds['max'][1]
                spanning=[c for c in table.cells if c[0]>=rx0-1 and c[2]<=rx1+1 and c[1] <= (lo+hi)/2 < c[3]]
                range_text=clean(' '.join(dict.fromkeys(cell_text(page,c) for c in spanning)))
                parts=re.split(r'\s+to\s+',range_text)
                if len(parts)==2:vals['min'],vals['max']=parts
            cond=vals.get('condition','')
            # Product group headings span limits too; inspect the full band only if no limits exist.
            if not any(vals.get(k) for k in ['min','typ','max']):
                band=cell_text(page,(bounds.get('condition',(sx1,0))[0],lo,table.bbox[2],hi))
                if PRODUCT_RE.search(band):active_products=products(band,defaults)
                continue
            ps=products(cond,active_products)
            for product in ps:
                records.append(row_record(product,vals,page_no,kind,section,notes))
    # Geometry can produce identical thin subrows; remove only identical rows within this table.
    seen=set(); result=[]
    for r in records:
        key=tuple(str(r[c]) for c in COLUMNS)
        if key not in seen:seen.add(key);result.append(r)
    return result

def frequency_table(page,table,defaults,page_no,section,notes):
    raw=table.extract()
    if not raw or 'Address Mode' not in clean(raw[0][0]):return []
    start=next((i for i,r in enumerate(raw) if any(clean(c)=='Aligned' for c in r)),None)
    if start is None:return []
    # Real merged cell geometry supplies instruction/opcode above each frequency column.
    rows=[]; heading=page.crop((0,max(0,table.bbox[1]-65),page.width,table.bbox[1])).extract_text() or ''
    mode=next((m for m in ['QPI_DTR','SPI_DTR','QPI','SPI'] if m in heading),'SPI')
    for idx in range(start+1,len(raw)):
        rawrow=raw[idx]; cycle=next((clean(v) for v in rawrow[:2] if v and re.fullmatch(r'\d+(?:[–-]\d+)?(?:\(\d+\))?',clean(v))), '')
        if not cycle:continue
        for j in range(2,len(rawrow)):
            val=clean(rawrow[j])
            if not val or val.lower()=='x':continue
            cell=table.rows[idx].cells[j]
            if not cell:continue
            x=(cell[0]+cell[2])/2
            def at(rowidx):
                yy=(table.rows[rowidx].bbox[1]+table.rows[rowidx].bbox[3])/2
                c=next((c for c in table.cells if c[0]<=x<c[2] and c[1]<=yy<c[3]),None)
                return cell_text(page,c) if c else ''
            param=at(0); align=at(start)
            for addressrow in [1,2]:
                opcode=at(addressrow)
                if not opcode or opcode.lower()=='x':continue
                condition=f'{mode}; {clean(raw[addressrow][0])}; opcode={opcode}; Dummy={cycle}; Starting Address: {align}'
                if '200' in val:condition+='; DS bit enabled'
                for product in defaults:
                    r=row_record(product,{'parameter':param+' maximum clock frequency','condition':condition,'max':re.sub(r'\(\d+\)','',val),'unit':'MHz'},page_no,'AC',section,notes)
                    r['description']='Frequency matrix; symbol not specified in source'
                    rows.append(r)
    return rows

def parse_pdf(path, explicit_pages=None):
    extracted=extract_tables(Path(path),list(ELECTRICAL_SECTION_LABELS),explicit_pages or [])
    rows=[]; review=[]; sources=[]
    by_page=defaultdict(list)
    for t in extracted:by_page[t.page_number].append(t)
    with pdfplumber.open(path) as pdf:
        first=' '.join((p.extract_text() or '') for p in pdf.pages[:2])
        family=re.search(r'W25Q((?:\d{2}/)+\d{2})RW',first,re.I)
        defaults=['W25Q'+x+'RW' for x in family.group(1).split('/')] if family else products(first,[])
        if not defaults:defaults=['未辨識產品']
        labels={};texts={};current='Electrical Characteristics'
        # Fast text-only scan back to chapter start, including explicit continuation pages.
        with fitz.open(path) as text_pdf:
            for idx in range(max(by_page,default=1)):
                text=text_pdf[idx].get_text(sort=True)
                heading_pattern=r'^\s*[1-9]\d*\.\d+\.?\s+([A-Z/].+)$'
                headings=re.findall(heading_pattern,text,re.M)
                early=re.findall(heading_pattern,text.split('Notes:')[0],re.M)
                if early:current=clean(early[-1])
                labels[idx+1]=current
                if headings:current=clean(headings[-1])
        for n,entries in by_page.items():
            p=pdf.pages[n-1]; txt=texts.get(n) or p.extract_text() or ''
            label=labels.get(n,entries[0].section or '')
            footer=re.search(r'-\s*(\d+)\s*-',txt)
            printed=int(footer.group(1)) if footer else n
            notes=txt[txt.find('Notes:'):] if 'Notes:' in txt else ''
            kind='DC' if 'DC ' in label else 'AC' if 'AC ' in label else 'Other'
            # Explicit-page mode: headings and table contents still inform classification.
            if label.startswith('Custom'):
                kind='DC' if 'Current' in txt or 'DC Electrical' in txt else 'AC' if any(s in txt for s in ['Timing','Clock','Time','AC Electrical']) else 'Other'
            tables=p.find_tables(PDF_TABLE_SETTINGS)
            sources.append({'pdf_page':n,'page':printed,'section':label,'text':txt})
            for i,t in enumerate(tables,1):
                above=p.crop((0,max(0,t.bbox[1]-40),p.width,t.bbox[1])).extract_text() or ''
                section=label+' | '+clean(above)
                got=standard_table(p,t,defaults,printed,kind,section,notes)
                if not got:got=frequency_table(p,t,defaults,printed,section,notes)
                if got:
                    for r in got:r['_pdf_page']=n;r['_table']=i
                    rows.extend(got)
                else:review.append({'page':printed,'pdf_page':n,'table':i,'section':label,'reason':'未能可靠轉成固定欄位，請核對原表。','raw':t.extract()})
            if not tables:
                for e in entries:review.append({'page':printed,'pdf_page':n,'table':e.table_index,'section':label,'reason':'備援擷取結果，需人工核對。','raw':e.dataframe.values.tolist()})
    if not rows:raise ValueError('未找到可標準化的規格。請確認 PDF 有文字內容，或指定規格頁的 PDF 實際頁碼。')
    for r in rows:
        r['_status']='uncompared'
        if not r['_notes']:
            label=r['_section'].split(' | ')[0]
            following=[s for s in sources if s['pdf_page']>=r['_pdf_page'] and s['section']==label and 'Notes:' in s['text']]
            if following:r['_notes']=following[0]['text'].split('Notes:',1)[1].split('Publication Release Date:')[0]
        r['_notes']=re.split(r'\n\s*\d+\.\d+\s+',r['_notes'])[0]
        if 'Factory Mode' in r['condition']:
            voltage=re.search(r'Factory mode.*?VCC range from ([0-9.]+V\s+to\s+[0-9.]+V)',r['_notes'],re.I|re.S)
            if voltage:r['condition']+='; VCC = '+clean(voltage.group(1))
    # Source keeps extracted page text; unsupported tables remain in Review.
    return {'rows':rows,'review':review,'sources':sources,'products':sorted({r['產品名稱'] for r in rows})}

def identity(r):
    return tuple(clean(r.get(k)).casefold() for k in ['產品名稱','spec type','symbol','parameter','condition'])

def compare_rows(new,old):
    counts=Counter(identity(r) for r in new); oldcounts=Counter(identity(r) for r in old)
    lookup={identity(r):r for r in old}; result=[]; seen=set()
    for source in new:
        r=dict(source);key=identity(r);seen.add(key);prev=lookup.get(key)
        r['_changes']={}
        if counts[key]>1 or oldcounts[key]>1:r['_status']='review'
        elif prev is None:r['_status']='added'
        else:
            for c in ['max','typ','min','unit','description']:
                if str(r.get(c,''))!=str(prev.get(c,'')):r['_changes'][c]={'old':prev.get(c,''),'new':r.get(c,'')}
            r['_status']='changed' if r['_changes'] else 'unchanged'
        result.append(r)
    result.extend(dict(r,_status='removed',_changes={}) for r in old if identity(r) not in seen)
    return result

def make_workbook(result):
    wb=Workbook();wb.remove(wb.active)
    def sheet(name,columns,rows):
        ws=wb.create_sheet(name);ws.append(columns)
        for row in rows:
            ws.append([row.get(k,'') for k in columns])
        for cells in ws:
            for c in cells:
                if isinstance(c.value,str):c.data_type='s'
                c.alignment=Alignment(vertical='top',wrap_text=True)
        for c in ws[1]:c.font=Font(color='FFFFFF',bold=True);c.fill=PatternFill('solid',fgColor='203E58')
        ws.freeze_panes='B2';ws.auto_filter.ref=ws.dimensions
        for col in ws.columns:
            letter=col[0].column_letter;header=col[0].value
            ws.column_dimensions[letter].width=45 if header in ['parameter','condition','description','text','raw'] else 19
        return ws
    current=[r for r in result['rows'] if r.get('_status')!='removed']
    ws=sheet('Specs',COLUMNS,current)
    for i,r in enumerate(current,2):
        color={'added':'DDF4E5','changed':'FFF0D1','review':'FADCE0'}.get(r.get('_status'))
        if color:
            for c in ws[i]:c.fill=PatternFill('solid',fgColor=color)
    changes=[]
    for r in result['rows']:
        status=r.get('_status')
        if status in ['changed','added','removed','review']:
            for field,delta in (r.get('_changes') or {'':{'old':'','new':''}}).items():
                changes.append(dict(r,status=status,field=field,old=delta['old'],new=delta['new']))
    sheet('Changes',['status']+COLUMNS+['field','old','new'],changes)
    reviews=[dict(r,raw='\n'.join(' | '.join(clean(c) for c in line) for line in r['raw'])) for r in result.get('review',[])]
    sheet('Review',['page','pdf_page','table','section','reason','raw'],reviews)
    sheet('Source',['page','pdf_page','section','text'],result.get('sources',[]))
    b=BytesIO();wb.save(b);return b.getvalue()
